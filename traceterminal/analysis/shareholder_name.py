"""Shareholder-classification Excel analyzer.

Search KSEI investor-classification files (DATE, SHARE CODE, ISSUER NAME,
then one column per investor classification and a TOTAL SCRIPLESS column)
by SHARE CODE.  Show the GOVERNMENT / POLITICAL PARTIES / STATE OWNED
COMPANY / CORPORATE / INDIVIDUAL / TOTAL SCRIPLESS breakdown across the
last 2/3/6/12 months, plus a table of which columns changed between the
earliest and the latest date in the selected period.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from traceterminal.analysis.shareholder5 import (
    NOT_IN_DOCUMENT,
    UNKNOWN,
    _clean,
    _date_from_filename,
    _date_key,
    _to_int,
)


# (label shown in the report, canonical key used in row data)
CLASSIFICATION_COLUMNS = [
    ("GOVERNMENT", "GOVERNMENT"),
    ("POLITICAL PARTIES", "POLITICAL_PARTIES"),
    ("STATE OWNED COMPANY", "STATE_OWNED_COMPANY"),
    ("CORPORATE", "CORPORATE"),
    ("INDIVIDUAL", "INDIVIDUAL"),
]
TOTAL_COLUMNS = [("TOTAL SCRIPLESS", "TOTAL_SCRIPLESS")]
ALL_COLUMNS = CLASSIFICATION_COLUMNS + TOTAL_COLUMNS


@dataclass(frozen=True)
class ClassificationRow:
    """One emiten row extracted from a shareholder-classification Excel file."""

    source_file: str
    report_date: str
    data: dict[str, Any]


def analyze_shareholder_name_documents(
    documents_dir: Path,
    share_code: str,
    period_months: int | None = None,
) -> dict[str, Any]:
    """Search classification Excel files by SHARE CODE and compare columns across dates."""
    target = share_code.strip().upper()
    if not target:
        raise ValueError("Kode emiten / SHARE CODE wajib diisi.")
    if not documents_dir.exists():
        raise FileNotFoundError(f"Folder dokumen tidak ditemukan: {documents_dir}")

    xlsx_files = sorted(
        documents_dir.glob("*.xlsx"),
        key=lambda path: _date_key(_date_from_filename(path)),
    )
    if not xlsx_files:
        return _empty_result(
            target,
            period_months,
            [f"Tidak ada file Excel di {documents_dir}."],
        )

    all_rows: list[ClassificationRow] = []
    errors: list[str] = []
    for xlsx_file in xlsx_files:
        try:
            all_rows.extend(_parse_rows(xlsx_file))
        except ImportError as exc:
            raise ImportError(
                "Library openpyxl belum tersedia. Jalankan `pip install -r requirements.txt`."
            ) from exc
        except Exception as exc:
            errors.append(f"{xlsx_file.name}: {exc}")

    matched = [row for row in all_rows if row.data.get("SHARE_CODE") == target]
    dates = sorted({row.report_date for row in matched}, key=_date_key)
    if period_months and len(dates) > period_months:
        selected = set(dates[-period_months:])
        matched = [row for row in matched if row.report_date in selected]

    aggregated = _aggregate_by_date(matched)
    rows = [dict(item.data) for item in aggregated]
    changes = _build_column_changes(aggregated)

    notes = [
        "Pencarian memakai Kode Emiten / SHARE CODE pada dokumen klasifikasi investor KSEI.",
        f"Kolom yang ditampilkan: {', '.join(label for label, _ in CLASSIFICATION_COLUMNS)} dan TOTAL SCRIPLESS.",
        "Perbandingan memakai tanggal terawal dan terbaru pada periode yang dipilih; "
        "tabel perubahan hanya menampilkan kolom yang nilainya berubah.",
    ]
    if period_months:
        notes.append(f"Hanya {period_months} bulan terakhir dari data yang tersedia yang dianalisa.")
    if errors:
        notes.append("Sebagian dokumen gagal dibaca: " + "; ".join(errors))
    if not rows:
        notes.append("Tidak ada data untuk Kode Emiten tersebut pada dokumen yang tersedia.")
    elif len(dates) < 2:
        notes.append("Perbandingan belum dapat dilakukan karena data tanggal pembanding belum tersedia.")

    return {
        "share_code": target,
        "period_months": period_months or "ALL",
        "rows": rows,
        "changes": changes,
        "documents_read": [xlsx.name for xlsx in xlsx_files],
        "total_rows_read": len(all_rows),
        "notes": notes,
    }


def _empty_result(
    share_code: str,
    period_months: int | None,
    notes: list[str],
) -> dict[str, Any]:
    return {
        "share_code": share_code,
        "period_months": period_months or "ALL",
        "rows": [],
        "changes": [],
        "documents_read": [],
        "total_rows_read": 0,
        "notes": notes,
    }


def _parse_rows(excel_file: Path) -> list[ClassificationRow]:
    """Read one classification Excel file and build one row per emiten.

    Data rows are laid out as DATE, SHARE CODE, ISSUER NAME, then one numeric
    column per investor classification ending with TOTAL SCRIPLESS.
    """
    try:
        import openpyxl  # type: ignore
    except ImportError:
        raise ImportError(
            "Library openpyxl belum tersedia. Jalankan `pip install -r requirements.txt`."
        ) from None

    rows: list[ClassificationRow] = []
    headers: list[str] | None = None
    index_map: dict[str, int] = {}

    wb = openpyxl.load_workbook(str(excel_file), read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        for raw_row in ws.iter_rows(values_only=True):
            cells = [_excel_cell(cell) for cell in raw_row]
            if not any(cells):
                continue

            if headers is None:
                if _is_header_row(cells):
                    headers = cells
                    index_map = _column_index_map(headers)
                continue

            date = _excel_cell(raw_row[0] if raw_row else None)
            report_date = _date_from_text(date) or _date_from_filename(excel_file)
            code = _table_cell(cells, 1).upper()
            issuer = _table_cell(cells, 2) or NOT_IN_DOCUMENT
            if not code or not _looks_like_share_code(code):
                continue

            data: dict[str, Any] = {
                "SHARE_CODE": code,
                "ISSUER_NAME": issuer,
                "REPORT_DATE": report_date,
            }
            for label, key in ALL_COLUMNS:
                index = index_map.get(key)
                value = _to_int(_table_cell(cells, index)) if index is not None else 0
                data[key] = value if isinstance(value, int) else 0
            rows.append(ClassificationRow(excel_file.name, report_date, data))
        return rows
    finally:
        wb.close()


def _excel_cell(value: Any) -> str:
    """Convert an Excel cell value to a clean string (dates to YYYY-MM-DD)."""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    return _clean(value)


def _table_cell(row: list[str], idx: int) -> str:
    return row[idx] if idx < len(row) else ""


def _is_header_row(cells: list[str]) -> bool:
    upper = [cell.upper() for cell in cells]
    return ("SHARE CODE" in upper or "SHARE_CODE" in upper) and "ISSUER NAME" in upper


def _column_index_map(headers: list[str]) -> dict[str, int]:
    """Map canonical column keys to their header index in the Excel file."""
    index_map: dict[str, int] = {}
    for index, label in enumerate(headers):
        key = re.sub(r"[^A-Z0-9]", "", label.upper())
        if key == "GOVERNMENT":
            index_map["GOVERNMENT"] = index
        elif key == "POLITICALPARTIES":
            index_map["POLITICAL_PARTIES"] = index
        elif key == "STATEOWNEDCOMPANY":
            index_map["STATE_OWNED_COMPANY"] = index
        elif key == "CORPORATE":
            index_map["CORPORATE"] = index
        elif key.startswith("INDIVIDUAL"):
            index_map["INDIVIDUAL"] = index
        elif key == "TOTALSCRIPLESS":
            index_map["TOTAL_SCRIPLESS"] = index
    return index_map


def _looks_like_share_code(value: Any) -> bool:
    return bool(re.fullmatch(r"[A-Z0-9]{4}", str(value or "").strip().upper()))


def _date_from_text(text: str) -> str | None:
    match = re.search(r"(\d{4})-(\d{1,2})-(\d{1,2})", str(text or ""))
    if match:
        year, month, day = map(int, match.groups())
        return f"{year:04d}-{month:02d}-{day:02d}"
    return None


def _aggregate_by_date(rows: list[ClassificationRow]) -> list[ClassificationRow]:
    """Merge rows sharing the same report date (and share code) by summing columns."""
    by_date: dict[str, ClassificationRow] = {}
    for row in sorted(rows, key=lambda item: _date_key(item.report_date)):
        existing = by_date.get(row.report_date)
        if existing is None:
            by_date[row.report_date] = row
            continue
        data = dict(existing.data)
        for _, key in ALL_COLUMNS:
            data[key] = int(existing.data.get(key) or 0) + int(row.data.get(key) or 0)
        if len(str(row.data.get("ISSUER_NAME", ""))) > len(str(existing.data.get("ISSUER_NAME", ""))):
            data["ISSUER_NAME"] = row.data.get("ISSUER_NAME")
        by_date[row.report_date] = ClassificationRow(
            existing.source_file,
            existing.report_date,
            data,
        )
    return [by_date[date] for date in sorted(by_date, key=_date_key)]


def _build_column_changes(rows: list[ClassificationRow]) -> list[dict[str, Any]]:
    """Compare earliest vs latest date and report columns whose value changed."""
    ordered = sorted(rows, key=lambda item: _date_key(item.report_date))
    if len(ordered) < 2:
        return []
    first = ordered[0]
    latest = ordered[-1]

    changes: list[dict[str, Any]] = []
    for label, key in ALL_COLUMNS:
        previous = int(first.data.get(key) or 0)
        current = int(latest.data.get(key) or 0)
        delta = current - previous
        if delta == 0:
            continue
        changes.append(
            {
                "KOLOM": label,
                "FROM_DATE": first.report_date,
                "TO_DATE": latest.report_date,
                "PREVIOUS_VALUE": previous,
                "LATEST_VALUE": current,
                "CHANGE_SHARES": delta,
                "ESTIMATED_LOTS": delta / 100,
                "STATUS": "Akumulasi" if delta > 0 else "Distribusi",
            }
        )
    return sorted(changes, key=lambda item: -abs(item.get("CHANGE_SHARES") or 0))