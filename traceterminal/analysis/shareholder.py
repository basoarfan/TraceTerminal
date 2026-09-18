"""Shareholder ownership Excel parser and analyzer."""

from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any


UNKNOWN = "Data tidak tersedia"
UNCONFIRMED = "Belum dapat dipastikan"
NOT_IN_DOCUMENT = "Tidak tercantum pada dokumen"

REQUIRED_COLUMNS = [
    "DATE",
    "SHARE_CODE",
    "ISSUER_NAME",
    "INVESTOR_NAME",
    "INVESTOR_CLASSIFICATION",
    "LOCAL_FOREIGN",
    "NATIONALITY",
    "DOMICILE",
    "HOLDINGS_SCRIPLESS",
    "HOLDINGS_SCRIP",
    "TOTAL_HOLDING_SHARES",
    "PERCENTAGE",
]


HEADER_ALIASES = {
    "date": "DATE",
    "tanggal": "DATE",
    "share code": "SHARE_CODE",
    "kode saham": "SHARE_CODE",
    "issuer name": "ISSUER_NAME",
    "nama emiten": "ISSUER_NAME",
    "investor name": "INVESTOR_NAME",
    "nama investor": "INVESTOR_NAME",
    "investor classification": "INVESTOR_CLASSIFICATION",
    "investor type": "INVESTOR_CLASSIFICATION",
    "klasifikasi investor": "INVESTOR_CLASSIFICATION",
    "local foreign": "LOCAL_FOREIGN",
    "local/foreign": "LOCAL_FOREIGN",
    "lokal asing": "LOCAL_FOREIGN",
    "nationality": "NATIONALITY",
    "kebangsaan": "NATIONALITY",
    "domicile": "DOMICILE",
    "domisili": "DOMICILE",
    "holdings scripless": "HOLDINGS_SCRIPLESS",
    "scripless": "HOLDINGS_SCRIPLESS",
    "holdings scrip": "HOLDINGS_SCRIP",
    "scrip": "HOLDINGS_SCRIP",
    "total holding shares": "TOTAL_HOLDING_SHARES",
    "total saham": "TOTAL_HOLDING_SHARES",
    "percentage": "PERCENTAGE",
    "persentase": "PERCENTAGE",
    "%": "PERCENTAGE",
}


@dataclass(frozen=True)
class ShareholderRow:
    """One ownership row extracted from source documents."""

    source_file: str
    data: dict[str, Any]


def analyze_shareholder_documents(
    documents_dir: Path,
    share_code: str,
    period_months: int | None = None,
) -> dict[str, Any]:
    """Read shareholder Excel files and compare ownership across available dates."""
    target = share_code.strip().upper()
    if not target:
        raise ValueError("Kode emiten wajib diisi.")

    if not documents_dir.exists():
        raise FileNotFoundError(f"Folder dokumen tidak ditemukan: {documents_dir}")

    xlsx_files = sorted(documents_dir.glob("*.xlsx"))
    if not xlsx_files:
        return {
            "share_code": target,
            "rows": [],
            "changes": [],
            "notes": [f"Tidak ada file Excel di {documents_dir}."],
        }

    rows: list[ShareholderRow] = []
    errors: list[str] = []
    for xlsx_file in xlsx_files:
        try:
            rows.extend(_parse_excel(xlsx_file))
        except ImportError as exc:
            raise ImportError(
                "Library openpyxl belum tersedia. Jalankan `pip install -r requirements.txt`."
            ) from exc
        except Exception as exc:  # Keep processing other documents.
            errors.append(f"{xlsx_file.name}: {exc}")

    filtered = [
        row for row in rows
        if str(row.data.get("SHARE_CODE", "")).upper() == target
    ]
    if period_months:
        all_dates = sorted(
            {
                str(row.data.get("DATE"))
                for row in rows
                if row.data.get("DATE") not in (None, UNKNOWN)
            },
            key=_date_key,
        )
        selected_dates = set(all_dates[-period_months:])
        filtered = [
            row for row in filtered
            if str(row.data.get("DATE")) in selected_dates
        ]
    filtered.sort(key=lambda row: (_date_key(row.data.get("DATE")), row.data.get("INVESTOR_NAME", "")))

    return {
        "share_code": target,
        "period_months": period_months or "ALL",
        "rows": [row.data | {"SOURCE_FILE": row.source_file} for row in filtered],
        "changes": _build_changes(filtered),
        "notes": _build_notes(filtered, errors),
        "documents_read": [xlsx.name for xlsx in xlsx_files],
        "total_rows_read": len(rows),
    }


def scan_shareholder_documents(
    documents_dir: Path,
    period_months: int,
) -> dict[str, Any]:
    """Scan every issuer and return only ownership movements in the period."""
    if period_months not in {2, 3, 6, 12}:
        raise ValueError("Periode scanner harus 2, 3, 6, atau 12 bulan.")

    if not documents_dir.exists():
        raise FileNotFoundError(f"Folder dokumen tidak ditemukan: {documents_dir}")

    xlsx_files = sorted(documents_dir.glob("*.xlsx"))
    if not xlsx_files:
        return {
            "period_months": period_months,
            "changes": [],
            "notes": [f"Tidak ada file Excel di {documents_dir}."],
            "documents_read": [],
            "total_rows_read": 0,
        }

    rows: list[ShareholderRow] = []
    errors: list[str] = []
    for xlsx_file in xlsx_files:
        try:
            rows.extend(_parse_excel(xlsx_file))
        except ImportError as exc:
            raise ImportError(
                "Library openpyxl belum tersedia. Jalankan `pip install -r requirements.txt`."
            ) from exc
        except Exception as exc:  # Keep processing other documents.
            errors.append(f"{xlsx_file.name}: {exc}")

    all_dates = sorted(
        {
            str(row.data.get("DATE"))
            for row in rows
            if row.data.get("DATE") not in (None, UNKNOWN)
        },
        key=_date_key,
    )
    selected_dates = set(all_dates[-period_months:])
    selected_rows = [
        row for row in rows
        if str(row.data.get("DATE")) in selected_dates
    ]

    rows_by_issuer: dict[str, list[ShareholderRow]] = {}
    for row in selected_rows:
        share_code = str(row.data.get("SHARE_CODE") or "").upper()
        if share_code:
            rows_by_issuer.setdefault(share_code, []).append(row)

    changes: list[dict[str, Any]] = []
    for share_code in sorted(rows_by_issuer):
        issuer_rows = rows_by_issuer[share_code]
        issuer_rows.sort(
            key=lambda row: (
                _date_key(row.data.get("DATE")),
                str(row.data.get("INVESTOR_NAME", "")),
            )
        )
        issuer_name = next(
            (
                row.data.get("ISSUER_NAME")
                for row in reversed(issuer_rows)
                if row.data.get("ISSUER_NAME") not in (None, "", NOT_IN_DOCUMENT)
            ),
            NOT_IN_DOCUMENT,
        )
        for change in _build_changes(issuer_rows):
            delta = _safe_int(change.get("CHANGE_SHARES"))
            if delta in (None, 0):
                continue
            changes.append(
                {
                    "SHARE_CODE": share_code,
                    "ISSUER_NAME": issuer_name,
                    **change,
                }
            )

    changes.sort(
        key=lambda change: (
            _date_key(change.get("FROM_DATE")),
            _date_key(change.get("TO_DATE")),
            str(change.get("SHARE_CODE", "")),
            str(change.get("INVESTOR_NAME", "")),
        )
    )

    notes = [
        "Scanner membaca seluruh kode emiten tanpa meminta kode saham.",
        "Hanya perubahan kepemilikan yang ditampilkan; posisi yang tidak berubah disembunyikan.",
    ]
    if 1 < len(selected_dates) < period_months:
        notes.append(
            f"Periode {period_months} bulan dipilih, tetapi hanya "
            f"{len(selected_dates)} tanggal laporan yang tersedia."
        )
    if len(selected_dates) < 2:
        notes.append(
            "Analisa belum bisa dilakukan karena data tanggal pembanding belum tersedia."
        )
    if errors:
        notes.append("Sebagian dokumen gagal dibaca: " + "; ".join(errors))

    return {
        "period_months": period_months,
        "dates_scanned": sorted(selected_dates, key=_date_key),
        "changes": changes,
        "notes": notes,
        "documents_read": [xlsx.name for xlsx in xlsx_files],
        "total_rows_read": len(rows),
        "total_issuers_scanned": len(rows_by_issuer),
        "total_changes": len(changes),
    }


def scan_shareholder_by_investor(
    documents_dir: Path,
    investor_name: str,
    period_months: int,
) -> dict[str, Any]:
    """Find one investor's ownership movements across every 1% issuer."""
    query = re.sub(r"\s+", " ", investor_name).strip()
    if not query:
        raise ValueError("Nama Pemegang Saham wajib diisi.")

    query_key = _investor_match_key(query)
    if query_key == UNKNOWN:
        raise ValueError("Nama Pemegang Saham tidak valid.")

    analysis = scan_shareholder_documents(documents_dir, period_months)
    matches = []
    for change in analysis.get("changes", []):
        investor_key = _investor_match_key(change.get("INVESTOR_NAME"))
        if query_key in investor_key or investor_key in query_key:
            matches.append(change)

    matched_names = sorted(
        {
            str(change.get("INVESTOR_NAME"))
            for change in matches
            if change.get("INVESTOR_NAME") not in (None, "", UNKNOWN)
        }
    )
    matched_issuers = {
        str(change.get("SHARE_CODE"))
        for change in matches
        if change.get("SHARE_CODE")
    }
    notes = list(analysis.get("notes", []))
    notes.append(
        "Pencarian nama tidak membedakan huruf besar/kecil dan mendukung nama sebagian."
    )
    if not matches:
        notes.append(
            f"Tidak ada perubahan kepemilikan untuk '{query}' pada periode terpilih."
        )

    return {
        **analysis,
        "investor_query": query,
        "matched_investor_names": matched_names,
        "changes": matches,
        "notes": notes,
        "total_issuers_matched": len(matched_issuers),
        "total_changes": len(matches),
    }


def _parse_excel(excel_file: Path) -> list[ShareholderRow]:
    """Read an Excel workbook and build ownership rows from its first sheet."""
    try:
        import openpyxl  # type: ignore
    except ImportError:
        raise ImportError(
            "Library openpyxl belum tersedia. Jalankan `pip install -r requirements.txt`."
        ) from None

    wb = openpyxl.load_workbook(str(excel_file), read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        rows: list[ShareholderRow] = []
        headers: list[str] | None = None
        for raw_row in ws.iter_rows(values_only=True):
            cells = [_excel_cell(cell) for cell in raw_row]
            if not any(cells):
                continue

            mapped = [_map_header(cell) for cell in cells]
            if "SHARE_CODE" in mapped and "INVESTOR_NAME" in mapped:
                headers = mapped
                continue

            if headers is None:
                continue

            data = _row_from_cells(headers, cells, _date_from_filename(excel_file))
            if data:
                rows.append(ShareholderRow(source_file=excel_file.name, data=data))
        return rows
    finally:
        wb.close()


def _excel_cell(value: Any) -> str:
    """Convert an Excel cell value to a clean string (dates to YYYY-MM-DD)."""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    return _clean_cell(value)


def _row_from_cells(
    headers: list[str],
    cells: list[str],
    default_date: str,
) -> dict[str, Any] | None:
    data = {column: NOT_IN_DOCUMENT for column in REQUIRED_COLUMNS}
    data["DATE"] = default_date

    for header, cell in zip(headers, cells):
        if header in REQUIRED_COLUMNS:
            data[header] = cell or NOT_IN_DOCUMENT

    if not _looks_like_share_code(data.get("SHARE_CODE")):
        return None

    data["SHARE_CODE"] = str(data["SHARE_CODE"]).upper()
    _normalize_numeric_fields(data)
    return data


def _build_changes(rows: list[ShareholderRow]) -> list[dict[str, Any]]:
    by_date: dict[str, dict[str, dict[str, Any]]] = {}
    for row in rows:
        investor = _investor_match_key(row.data.get("INVESTOR_NAME"))
        date = str(row.data.get("DATE") or UNKNOWN)
        date_rows = by_date.setdefault(date, {})
        existing_key = _find_similar_investor_key(date_rows, investor)
        if existing_key:
            date_rows[existing_key] = _merge_investor_rows(date_rows[existing_key], row.data)
        else:
            date_rows[investor] = row.data

    changes: list[dict[str, Any]] = []
    ordered_dates = sorted(by_date, key=_date_key)
    if len(ordered_dates) < 2:
        for date in ordered_dates:
            for investor_key, row in sorted(by_date[date].items()):
                changes.append(
                    {
                        "INVESTOR_NAME": row.get("INVESTOR_NAME", investor_key),
                        "STATUS": UNCONFIRMED,
                        "DETAIL": "Data pembanding tanggal sebelumnya belum tersedia.",
                    }
                )
        return changes

    for previous_date, current_date in zip(ordered_dates, ordered_dates[1:]):
        previous_rows = by_date[previous_date]
        current_rows = by_date[current_date]
        investor_keys = sorted(set(previous_rows) | set(current_rows))

        for investor_key in investor_keys:
            previous = previous_rows.get(investor_key)
            current = current_rows.get(investor_key)

            if previous is None and current is not None:
                current_total = _safe_int(current.get("TOTAL_HOLDING_SHARES"))
                changes.append(
                    {
                        "INVESTOR_NAME": current.get("INVESTOR_NAME", investor_key),
                        "FROM_DATE": previous_date,
                        "TO_DATE": current_date,
                        "STATUS": "Masuk daftar >1%",
                        "CHANGE_SHARES": current_total,
                        "ESTIMATED_LOTS": current_total / 100 if current_total is not None else None,
                        "STRENGTH": "kuat",
                        "PREVIOUS_TOTAL": 0,
                        "CURRENT_TOTAL": current.get("TOTAL_HOLDING_SHARES", UNKNOWN),
                    }
                )
                continue

            if previous is not None and current is None:
                previous_total = _safe_int(previous.get("TOTAL_HOLDING_SHARES"))
                delta = -previous_total if previous_total is not None else None
                changes.append(
                    {
                        "INVESTOR_NAME": previous.get("INVESTOR_NAME", investor_key),
                        "FROM_DATE": previous_date,
                        "TO_DATE": current_date,
                        "STATUS": "Keluar daftar >1%",
                        "CHANGE_SHARES": delta,
                        "ESTIMATED_LOTS": delta / 100 if delta is not None else None,
                        "STRENGTH": "kuat",
                        "PREVIOUS_TOTAL": previous.get("TOTAL_HOLDING_SHARES", UNKNOWN),
                        "CURRENT_TOTAL": 0,
                    }
                )
                continue

            if previous is None or current is None:
                continue

            previous_total = _safe_int(previous.get("TOTAL_HOLDING_SHARES"))
            current_total = _safe_int(current.get("TOTAL_HOLDING_SHARES"))
            if previous_total is None or current_total is None:
                status = UNCONFIRMED
                delta = None
                lots = None
                strength = UNCONFIRMED
            else:
                delta = current_total - previous_total
                lots = delta / 100
                status = "Akumulasi" if delta > 0 else "Distribusi" if delta < 0 else "Netral"
                strength = _movement_strength(previous_total, delta)

            changes.append(
                {
                    "INVESTOR_NAME": current.get("INVESTOR_NAME", investor_key),
                    "FROM_DATE": previous.get("DATE", UNKNOWN),
                    "TO_DATE": current.get("DATE", UNKNOWN),
                    "STATUS": status,
                    "CHANGE_SHARES": delta,
                    "ESTIMATED_LOTS": lots,
                    "STRENGTH": strength,
                    "PREVIOUS_TOTAL": previous.get("TOTAL_HOLDING_SHARES", UNKNOWN),
                    "CURRENT_TOTAL": current.get("TOTAL_HOLDING_SHARES", UNKNOWN),
                }
            )

    return changes


def _build_notes(rows: list[ShareholderRow], errors: list[str]) -> list[str]:
    notes = [
        "Analisa hanya menggunakan data yang berhasil dibaca dari dokumen Excel.",
        f"Data broker: {NOT_IN_DOCUMENT}.",
    ]
    if not rows:
        notes.append("Tidak ada data untuk kode emiten tersebut pada dokumen yang tersedia.")
    elif len({_date_key(row.data.get("DATE")) for row in rows}) < 2:
        notes.append("Analisa akumulasi/distribusi belum bisa dipastikan karena data tanggal pembanding belum tersedia.")
    if errors:
        notes.append("Sebagian dokumen gagal dibaca: " + "; ".join(errors))
    return notes


def _movement_strength(previous_total: int, delta: int) -> str:
    if delta == 0 or previous_total == 0:
        return "netral"
    pct = abs(delta) / previous_total * 100
    if pct >= 5:
        return "kuat"
    if pct >= 1:
        return "signifikan"
    return "kecil"


def _investor_match_key(value: Any) -> str:
    text = str(value or "").upper()
    text = re.sub(r"[^A-Z0-9 ]+", " ", text)
    text = _strip_investor_suffix(text)
    legal_words = {
        "PT",
        "TBK",
        "PTE",
        "LTD",
        "LIMITED",
        "CORP",
        "CORPORATION",
        "COMPANY",
    }
    tokens = [
        token for token in re.sub(r"\s+", " ", text).strip().split()
        if token not in legal_words
    ]
    return " ".join(tokens) or UNKNOWN


def _strip_investor_suffix(text: str) -> str:
    suffix_patterns = [
        r"\bINVESTMENT\s+ADVISORS?\s+[AF]\s+[A-Z ]+$",
        r"\bPRIVATE\s+EQUITY\s+[AF]\s+[A-Z ]+$",
        r"\bFINANCIAL\s+INSTITUTION(?:AL)?\s+[AF]\s+[A-Z ]+$",
        r"\bCORPORATE\s+[AF]\s+[A-Z ]+$",
        r"\bINDIVIDUAL\s+[AF]\s+[A-Z ]+$",
        r"\b[A-Z ]+\s+[AF]\s+(?:SINGAPORE|INDONESIA|MALAYSIA|HONG KONG|JAPAN|CHINA|USA|UNITED STATES|BRITISH VIRGIN ISLANDS)$",
    ]
    cleaned = re.sub(r"\s+", " ", text).strip()
    for pattern in suffix_patterns:
        cleaned = re.sub(pattern, "", cleaned).strip()
    return cleaned


def _merge_investor_rows(
    existing: dict[str, Any],
    new: dict[str, Any],
) -> dict[str, Any]:
    merged = dict(existing)
    for key in ("HOLDINGS_SCRIPLESS", "HOLDINGS_SCRIP", "TOTAL_HOLDING_SHARES"):
        left = _safe_int(existing.get(key))
        right = _safe_int(new.get(key))
        if left is not None and right is not None:
            merged[key] = left + right
    existing_name = str(existing.get("INVESTOR_NAME") or "")
    new_name = str(new.get("INVESTOR_NAME") or "")
    if len(new_name) > len(existing_name):
        merged["INVESTOR_NAME"] = new_name
    return merged


def _find_similar_investor_key(
    date_rows: dict[str, dict[str, Any]],
    investor_key: str,
) -> str | None:
    for existing_key in date_rows:
        if existing_key == investor_key:
            return existing_key
        if _is_similar_investor_key(existing_key, investor_key):
            return existing_key
    return None


def _is_similar_investor_key(left: str, right: str) -> bool:
    if not left or not right:
        return False
    left_tokens = left.split()
    right_tokens = right.split()
    shorter, longer = (
        (left_tokens, right_tokens)
        if len(left_tokens) <= len(right_tokens)
        else (right_tokens, left_tokens)
    )
    if len(shorter) < 2:
        return False
    return shorter == longer[:len(shorter)]


def _clean_cell(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").replace("\n", " ")).strip()


def _map_header(header: str) -> str:
    key = re.sub(r"[^a-z% ]+", " ", header.lower())
    key = re.sub(r"\s+", " ", key).strip()
    return HEADER_ALIASES.get(key, header.upper().strip())


def _looks_like_share_code(value: Any) -> bool:
    return bool(re.fullmatch(r"[A-Z0-9]{4}", str(value or "").strip().upper()))


def _normalize_numeric_fields(data: dict[str, Any]) -> None:
    for key in ("HOLDINGS_SCRIPLESS", "HOLDINGS_SCRIP", "TOTAL_HOLDING_SHARES"):
        data[key] = _to_int(data.get(key))
    data["PERCENTAGE"] = _to_float(data.get("PERCENTAGE"))


def _to_int(value: Any) -> int | str:
    text = str(value or "").strip()
    cleaned = re.sub(r"[^\d-]", "", text)
    if not cleaned:
        return NOT_IN_DOCUMENT
    return int(cleaned)


def _safe_int(value: Any) -> int | None:
    if isinstance(value, int):
        return value
    parsed = _to_int(value)
    return parsed if isinstance(parsed, int) else None


def _to_float(value: Any) -> float | str:
    text = str(value or "").strip().replace("%", "")
    text = re.sub(r"[^\d,.-]", "", text)
    if not text:
        return NOT_IN_DOCUMENT
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return NOT_IN_DOCUMENT


def _date_from_filename(path: Path) -> str:
    return _date_from_text(path.stem) or UNKNOWN


def _date_from_text(text: str) -> str | None:
    month_map = {
        "jan": 1, "january": 1, "januari": 1,
        "feb": 2, "february": 2, "februari": 2,
        "mar": 3, "march": 3, "maret": 3,
        "apr": 4, "april": 4,
        "may": 5, "mei": 5,
        "jun": 6, "june": 6, "juni": 6,
        "jul": 7, "july": 7, "juli": 7,
        "aug": 8, "august": 8, "agu": 8, "agustus": 8,
        "sep": 9, "sept": 9, "september": 9,
        "oct": 10, "october": 10, "okt": 10, "oktober": 10,
        "nov": 11, "november": 11,
        "dec": 12, "december": 12, "des": 12, "desember": 12,
    }
    match = re.search(r"(\d{4})[-/](\d{1,2})[-/](\d{1,2})", text)
    if match:
        year, month, day = map(int, match.groups())
        return f"{year:04d}-{month:02d}-{day:02d}"

    match = re.search(r"(\d{1,2})[-\s/]([A-Za-z]+)[-\s/](\d{4})", text)
    if match:
        day = int(match.group(1))
        month = month_map.get(match.group(2).lower())
        year = int(match.group(3))
        if month:
            return f"{year:04d}-{month:02d}-{day:02d}"

    match = re.search(r"(\d{1,2})[-/](\d{1,2})[-/](\d{4})", text)
    if match:
        day, month, year = map(int, match.groups())
        return f"{year:04d}-{month:02d}-{day:02d}"
    return None


def _date_key(value: Any) -> datetime:
    try:
        return datetime.strptime(str(value), "%Y-%m-%d")
    except ValueError:
        return datetime.min
